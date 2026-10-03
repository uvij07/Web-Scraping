import requests
import pandas as pd
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from openpyxl import load_workbook

books_info = []

for x in range(1, 51):
    website = f"https://books.toscrape.com/catalogue/page-{x}.html"

    r = requests.get(website)
    soup = BeautifulSoup(r.content, "html.parser")

    book = soup.find_all(
        "li",
        class_="col-xs-6 col-sm-4 col-md-3 col-lg-3"
    )

    for item in book:
        name_tag = item.find("h3").find("a")

        book_name = name_tag.text.strip()
        book_link = urljoin(website, name_tag["href"])
        price = item.find("p", class_="price_color").text.strip()

        all_products = {
            "Name": book_name,
            "Link": book_link,
            "Price": price
        }

        books_info.append(all_products)

df = pd.DataFrame(books_info)

file_name = "Books_info.xlsx"
df.to_excel(file_name, index=False)

# Open Excel file
wb = load_workbook(file_name)
ws = wb.active

# Make Link column clickable
for row in range(2, ws.max_row + 1):
    cell = ws[f"B{row}"]
    cell.hyperlink = cell.value
    cell.style = "Hyperlink"

wb.save(file_name)